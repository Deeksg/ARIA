"""
backend/ingestion/pipeline.py
=============================
Ingestion pipeline for ARIA — ingests resolved historical tickets (clean CSV)
and support documentation (.txt, .docx, .pdf) into ChromaDB vector store.
"""

import os
import sys
import time
import json
import hashlib
import logging
import argparse
import random
from pathlib import Path
from datetime import datetime, timezone
from typing import Union
import pandas as pd
from dotenv import load_dotenv

# ── ENVIRONMENT & VECTOR STORE IMPORTS ────────────────────────────────────────
# load_dotenv() must be called before project imports so environment variables are loaded
load_dotenv()

from backend.pipeline.vector_store import (
    PROJECT_ROOT,
    COLLECTION_NAME,
    UNCLASSIFIED,
    upsert_records,
    add_resolved_ticket,
    build_ticket_metadata,
    delete_by_source_file,
    search_similar,
    get_collection,
    get_collection_stats,
    reset_collection,
)

logger = logging.getLogger(__name__)

# ── CONSTANTS ─────────────────────────────────────────────────────────────────
# Default path to clean historical resolved tickets dataset (never raw)
DEFAULT_RESOLVED_CSV = PROJECT_ROOT / "data/synthetic/resolved_tickets_clean.csv"

# Supported documentation file extensions
SUPPORTED_EXTENSIONS = {".txt", ".docx", ".pdf"}

# Chunking configuration:
# 150 words with 30 words overlap, NOT 512 tokens, because all-MiniLM-L6-v2
# silently truncates input beyond 256 word pieces.
CHUNK_WORDS = 150
CHUNK_OVERLAP_WORDS = 30

REQUIRED_CSV_COLUMNS = [
    "ticket_id",
    "module",
    "issue_type",
    "problem",
    "error_code",
    "root_cause",
    "fix",
    "tcodes_used",
    "resolution_time_minutes",
    "escalated",
    "created_date",
    "resolved_date",
    "ticket_type",
]


# ── CSV INGESTION FUNCTIONS ───────────────────────────────────────────────────
def load_resolved_csv(csv_path: Union[str, Path]) -> tuple[list[dict], dict]:
    """
    Loads resolved tickets CSV, validates schema, and returns cleaned records.

    Args:
        csv_path: Path to resolved tickets clean CSV.

    Returns:
        tuple[list[dict], dict]: List of valid record dicts and summary stats.

    Raises:
        FileNotFoundError: If csv_path does not exist.
        ValueError: If any required column is missing.
    """
    path = Path(csv_path)
    if not path.exists():
        raise FileNotFoundError(f"Resolved tickets CSV not found at: {path}")

    df = pd.read_csv(path, dtype=str, keep_default_na=False)

    missing = [c for c in REQUIRED_CSV_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"Missing required columns in CSV: {missing}")

    # Strip surrounding whitespace from all string columns
    for col in df.columns:
        df[col] = df[col].astype(str).str.strip()

    total_rows = len(df)
    valid_records = []
    skipped_blank = 0

    for _, row in df.iterrows():
        t_id = row["ticket_id"]
        prob = row["problem"]

        # Skip rows with blank ticket_id or problem text
        if not t_id or not prob:
            skipped_blank += 1
            continue

        raw_min = row.get("resolution_time_minutes", "")
        try:
            res_min = int(float(raw_min)) if raw_min != "" else -1
        except (ValueError, TypeError):
            res_min = -1

        rec = {
            "ticket_id": t_id,
            "module": row.get("module", ""),
            "issue_type": row.get("issue_type", ""),
            "problem": prob,
            "error_code": row.get("error_code", ""),
            "root_cause": row.get("root_cause", ""),
            "fix": row.get("fix", ""),
            "tcodes_used": row.get("tcodes_used", ""),
            "resolution_time_minutes": res_min,
            "escalated": row.get("escalated", ""),
            "created_date": row.get("created_date", ""),
            "resolved_date": row.get("resolved_date", ""),
            "ticket_type": row.get("ticket_type", ""),
        }
        valid_records.append(rec)

    summary = {
        "total_rows": total_rows,
        "valid_records": len(valid_records),
        "skipped_blank": skipped_blank,
    }
    return valid_records, summary


def ingest_csv(
    csv_path: Union[str, Path] = DEFAULT_RESOLVED_CSV,
    batch_size: int = 500,
    force: bool = False,
) -> dict:
    """
    Ingests resolved tickets CSV into ChromaDB vector store.

    Embeds ONLY the problem text (Decision 2). All other columns are preserved
    in metadata. Uses skip_existing=not force for idempotency.

    Args:
        csv_path: Path to CSV dataset.
        batch_size: Upsert batch size.
        force: If True, overwrites existing records instead of skipping.

    Returns:
        dict: Ingestion execution summary.
    """
    start_time = time.time()
    csv_file = Path(csv_path)
    records, load_summary = load_resolved_csv(csv_file)

    ids = [r["ticket_id"] for r in records]
    texts = [r["problem"] for r in records]
    metadatas = []
    source_filename = csv_file.name

    for r in records:
        meta = build_ticket_metadata(r)
        meta["source_file"] = source_filename
        metadatas.append(meta)

    # Upsert with skip_existing = not force
    upsert_res = upsert_records(
        ids=ids,
        texts=texts,
        metadatas=metadatas,
        batch_size=batch_size,
        skip_existing=(not force),
    )

    elapsed = round(time.time() - start_time, 2)
    collection = get_collection()
    count_after = collection.count()

    summary = {
        "rows_read": load_summary["total_rows"],
        "skipped_blank": load_summary["skipped_blank"],
        "already_present": upsert_res["skipped_existing"],
        "newly_added": upsert_res["upserted"],
        "elapsed_seconds": elapsed,
        "collection_count_after": count_after,
    }

    print("=" * 60)
    print("CSV INGESTION COMPLETE")
    print("=" * 60)
    print(f"Source file:            {source_filename}")
    print(f"Rows read:              {summary['rows_read']:,}")
    print(f"Skipped blank:          {summary['skipped_blank']:,}")
    print(f"Already present:        {summary['already_present']:,}")
    print(f"Newly added:            {summary['newly_added']:,}")
    print(f"Elapsed time:           {summary['elapsed_seconds']}s")
    print(f"Total collection count: {summary['collection_count_after']:,}")
    print("=" * 60)

    return summary


# ── DOCUMENT PROCESSING & CHUNKING ────────────────────────────────────────────
def read_document(path: Union[str, Path]) -> str:
    """
    Extracts text from supported document files (.txt, .docx, .pdf).

    Args:
        path: Path to document file.

    Returns:
        str: Extracted text content.

    Raises:
        ValueError: If file extension is unsupported.
    """
    p = Path(path)
    ext = p.suffix.lower()

    if ext == ".txt":
        return p.read_text(encoding="utf-8", errors="replace")

    elif ext == ".docx":
        import docx
        doc = docx.Document(p)
        text_parts = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    c_text = cell.text.strip()
                    if c_text:
                        text_parts.append(c_text)
        return "\n".join(text_parts)

    elif ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(p)
        pages_text = []
        for i, page in enumerate(reader.pages):
            try:
                t = page.extract_text()
                if t and t.strip():
                    pages_text.append(t.strip())
            except Exception as e:
                logger.warning(f"Failed to extract page {i} of {p}: {e}")
        return "\n".join(pages_text)

    else:
        raise ValueError(f"Unsupported file extension: {ext}")


def chunk_text(
    text: str,
    chunk_words: int = CHUNK_WORDS,
    overlap_words: int = CHUNK_OVERLAP_WORDS,
) -> list[str]:
    """
    Splits text into overlapping word chunks.

    WHY 150 WORDS WITH 30 WORDS OVERLAP:
    Document chunks are 150 words with 30 words overlap, NOT 512 tokens,
    because all-MiniLM-L6-v2 silently truncates input beyond 256 word pieces.
    150 English words produce ~200 wordpieces, leaving comfortable headroom
    below the 256 token limit while preserving semantic context.

    Args:
        text: Input text string.
        chunk_words: Words per chunk.
        overlap_words: Overlapping words between consecutive chunks.

    Returns:
        list[str]: List of chunked text strings.

    Raises:
        ValueError: If overlap_words >= chunk_words.
    """
    if overlap_words >= chunk_words:
        raise ValueError(
            f"overlap_words ({overlap_words}) must be strictly less than "
            f"chunk_words ({chunk_words})"
        )

    words = text.split()
    if not words:
        return []

    start = 0
    chunks = []
    step = chunk_words - overlap_words

    while start < len(words):
        end = min(start + chunk_words, len(words))
        chunk = " ".join(words[start:end])
        chunks.append(chunk)
        if end == len(words):
            break
        start += step

    return chunks


def ingest_files(
    folder: Union[str, Path],
    module: str | None = None,
    issue_type: str | None = None,
) -> dict:
    """
    Walks a folder recursively and ingests supported documentation into ChromaDB.

    For each file:
    - Reads and chunks text.
    - Computes sha1 hash.
    - Generates chunk ID: 'doc-{sha1[:12]}-{i:04d}'.
    - Deletes previous chunks for this source_file to avoid stale chunks.
    - Upserts chunks with skip_existing=False.

    Args:
        folder: Folder path to walk.
        module: Optional SAP module metadata tag.
        issue_type: Optional issue type metadata tag.

    Returns:
        dict: Ingestion summary.
    """
    folder_path = Path(folder)
    if not folder_path.exists():
        raise FileNotFoundError(f"Folder not found: {folder_path}")

    files_found = 0
    files_ingested = 0
    skipped_empty = []
    skipped_unsupported = []
    total_chunks_added = 0

    for root, _, files in os.walk(folder_path):
        for f in files:
            files_found += 1
            f_path = Path(root) / f
            ext = f_path.suffix.lower()

            if ext not in SUPPORTED_EXTENSIONS:
                skipped_unsupported.append(str(f_path.relative_to(folder_path)))
                continue

            # Read document text
            try:
                raw_text = read_document(f_path)
            except Exception as e:
                logger.warning(f"Error reading {f_path}: {e}")
                skipped_empty.append(str(f_path.relative_to(folder_path)))
                continue

            if not raw_text or not raw_text.strip():
                logger.warning(f"No extractable text in {f_path}; skipping.")
                skipped_empty.append(str(f_path.relative_to(folder_path)))
                continue

            chunks = chunk_text(raw_text)
            if not chunks:
                skipped_empty.append(str(f_path.relative_to(folder_path)))
                continue

            # Compute sha1 hash of raw file bytes
            file_bytes = f_path.read_bytes()
            sha1 = hashlib.sha1(file_bytes).hexdigest()

            # Store relative source path with forward slashes
            rel_source_file = f_path.relative_to(folder_path).as_posix()

            # Delete old chunks for this source_file so edits leave no stale chunks
            delete_by_source_file(rel_source_file)

            now_iso = datetime.now(timezone.utc).isoformat()
            chunk_ids = []
            chunk_texts = []
            chunk_metas = []

            for i, chunk in enumerate(chunks):
                c_id = f"doc-{sha1[:12]}-{i:04d}"
                c_meta = {
                    "source_type": "document_chunk",
                    "source_file": rel_source_file,
                    "source_hash": sha1,
                    "chunk_index": i,
                    "total_chunks": len(chunks),
                    "module": module if module else UNCLASSIFIED,
                    "issue_type": issue_type if issue_type else UNCLASSIFIED,
                    "ingested_at": now_iso,
                }
                chunk_ids.append(c_id)
                chunk_texts.append(chunk)
                chunk_metas.append(c_meta)

            upsert_records(
                ids=chunk_ids,
                texts=chunk_texts,
                metadatas=chunk_metas,
                skip_existing=False,
            )
            files_ingested += 1
            total_chunks_added += len(chunks)

    summary = {
        "files_found": files_found,
        "files_ingested": files_ingested,
        "skipped_empty": skipped_empty,
        "skipped_unsupported": skipped_unsupported,
        "chunks_added": total_chunks_added,
    }

    print("=" * 60)
    print("FILE INGESTION COMPLETE")
    print("=" * 60)
    print(f"Files found:       {files_found}")
    print(f"Files ingested:    {files_ingested}")
    print(f"Chunks added:      {total_chunks_added}")
    print(f"Skipped empty:     {len(skipped_empty)}")
    print(f"Unsupported files: {len(skipped_unsupported)}")
    print("=" * 60)

    return summary


# ── VERIFICATION & DIAGNOSTIC BENCHMARKS ───────────────────────────────────────
def run_verify(sample_size: int = 300, k: int = 5) -> None:
    """
    Runs diagnostic and quality verification checks on the vector store:
    a) Detailed collection stats.
    b) 5 messy hand-written test queries.
    c) Leave-one-out retrieval benchmark over sample_size records.
    d) Caveat note regarding synthetic duplicates.
    """
    # a) Detailed collection stats
    print("\n" + "=" * 60)
    print("1. COLLECTION STATISTICS (DETAILED)")
    print("=" * 60)
    stats = get_collection_stats(detailed=True)
    for k_name, val in stats.items():
        if isinstance(val, dict):
            print(f"{k_name}:")
            for sub_k, sub_v in sorted(val.items()):
                print(f"  {sub_k}: {sub_v}")
        else:
            print(f"{k_name}: {val}")

    collection = get_collection()
    total_count = collection.count()
    if total_count == 0:
        print("\nCollection is empty. Ingest data first to run retrieval checks.")
        return

    # b) 5 messy hand-written test queries
    print("\n" + "=" * 60)
    print("2. MESSY QUERY RETRIEVAL SAMPLES")
    print("=" * 60)
    messy_queries = [
        "invoice not going through since yesterday please help",
        "cannot post goods receipt error M7001 missing data",
        "employee payroll calculation shows incorrect tax deduction",
        "user unable to login authorization failed su01 role missing",
        "sap gui crash during month end closing transaction",
    ]

    for q in messy_queries:
        print(f"\nQuery: '{q}'")
        results = search_similar(q, n_results=3)
        for rank, r in enumerate(results, 1):
            meta = r.get("metadata", {})
            mod = meta.get("module", "N/A")
            itype = meta.get("issue_type", "N/A")
            prob = r.get("text", "")[:70]
            root = meta.get("root_cause", "")[:60]
            sim = r.get("similarity", 0.0)
            print(f"  [{rank}] Sim={sim:.3f} | Mod={mod} | Type={itype}")
            print(f"      Problem: {prob}...")
            print(f"      Cause:   {root}...")

    # c) Leave-one-out self-check
    print("\n" + "=" * 60)
    print("3. LEAVE-ONE-OUT RETRIEVAL BENCHMARK")
    print("=" * 60)

    n_sample = min(sample_size, total_count)
    all_data = collection.get(
        limit=total_count,
        where={"source_type": "resolved_ticket"},
        include=["embeddings", "metadatas"],
    )

    all_ids = all_data.get("ids", [])
    all_embeddings = all_data.get("embeddings", [])
    all_metadatas = all_data.get("metadatas", [])

    if len(all_ids) == 0:
        print("No resolved_ticket records available for leave-one-out test.")
        return

    # Seed random sampling for deterministic verification
    rng = random.Random(42)
    sample_indices = rng.sample(range(len(all_ids)), min(n_sample, len(all_ids)))

    same_module_counts = []
    same_issue_counts = []
    same_both_counts = []
    top1_module_matches = 0
    top1_issue_matches = 0

    for idx in sample_indices:
        t_id = all_ids[idx]
        t_meta = all_metadatas[idx]
        t_emb = all_embeddings[idx]

        # Chroma may return embeddings as numpy arrays: never use truthiness (if t_emb)
        if t_emb is None or len(t_emb) == 0:
            continue

        emb_list = t_emb.tolist() if hasattr(t_emb, "tolist") else list(t_emb)

        # Query using stored embedding with k+1 results
        q_res = collection.query(
            query_embeddings=[emb_list],
            n_results=k + 1,
            include=["metadatas"],
        )

        res_ids = q_res.get("ids", [[]])[0]
        res_metas = q_res.get("metadatas", [[]])[0]

        # Drop the record itself by id, keeping the next k
        filtered_metas = [m for r_id, m in zip(res_ids, res_metas) if r_id != t_id][:k]
        if not filtered_metas:
            continue

        target_mod = t_meta.get("module")
        target_issue = t_meta.get("issue_type")

        mod_matches = sum(1 for m in filtered_metas if m.get("module") == target_mod)
        issue_matches = sum(1 for m in filtered_metas if m.get("issue_type") == target_issue)
        both_matches = sum(
            1 for m in filtered_metas
            if m.get("module") == target_mod and m.get("issue_type") == target_issue
        )

        same_module_counts.append(mod_matches / len(filtered_metas))
        same_issue_counts.append(issue_matches / len(filtered_metas))
        same_both_counts.append(both_matches / len(filtered_metas))

        if filtered_metas[0].get("module") == target_mod:
            top1_module_matches += 1
        if filtered_metas[0].get("issue_type") == target_issue:
            top1_issue_matches += 1

    total_tested = len(same_module_counts)
    if total_tested > 0:
        mean_mod = sum(same_module_counts) / total_tested
        mean_issue = sum(same_issue_counts) / total_tested
        mean_both = sum(same_both_counts) / total_tested
        top1_mod_rate = top1_module_matches / total_tested
        top1_issue_rate = top1_issue_matches / total_tested

        print(f"Sample size tested:          {total_tested}")
        print(f"Mean top-{k} same module:      {mean_mod*100:.1f}%")
        print(f"Mean top-{k} same issue_type:  {mean_issue*100:.1f}%")
        print(f"Mean top-{k} same both:        {mean_both*100:.1f}%")
        print(f"Top-1 same module rate:      {top1_mod_rate*100:.1f}%")
        print(f"Top-1 same issue_type rate:  {top1_issue_rate*100:.1f}%")

    # d) Caveat note
    print("\nNOTE:")
    print("These numbers are inflated: synthetic tickets are paraphrases of the same")
    print("seeds, so near-duplicates exist. They show the retrieval machinery works,")
    print("not real-world accuracy.")
    print("=" * 60)


# ── CLI ENTRY POINT ───────────────────────────────────────────────────────────
def main():
    """
    CLI command dispatcher for ingestion pipeline:
      python -m backend.ingestion.pipeline csv [--path PATH] [--force]
      python -m backend.ingestion.pipeline files FOLDER [--module M] [--issue-type T]
      python -m backend.ingestion.pipeline stats
      python -m backend.ingestion.pipeline verify
      python -m backend.ingestion.pipeline reset --yes
    """
    parser = argparse.ArgumentParser(
        description="ARIA Vector Store Ingestion Pipeline"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    # Subcommand: csv
    csv_parser = subparsers.add_parser("csv", help="Ingest resolved tickets CSV")
    csv_parser.add_argument(
        "--path",
        default=str(DEFAULT_RESOLVED_CSV),
        help=f"Path to CSV (default: {DEFAULT_RESOLVED_CSV})",
    )
    csv_parser.add_argument(
        "--force",
        action="store_true",
        help="Overwrite existing ticket IDs instead of skipping",
    )

    # Subcommand: files
    files_parser = subparsers.add_parser("files", help="Ingest documentation folder")
    files_parser.add_argument("folder", help="Path to folder containing documents")
    files_parser.add_argument("--module", default=None, help="SAP module metadata tag")
    files_parser.add_argument("--issue-type", default=None, help="Issue type metadata tag")

    # Subcommand: stats
    subparsers.add_parser("stats", help="Print vector store collection statistics")

    # Subcommand: verify
    verify_parser = subparsers.add_parser("verify", help="Run retrieval verification checks")
    verify_parser.add_argument("--sample-size", type=int, default=300, help="Benchmark sample size")
    verify_parser.add_argument("--k", type=int, default=5, help="Number of neighbors k")

    # Subcommand: reset
    reset_parser = subparsers.add_parser("reset", help="Reset and clear vector collection")
    reset_parser.add_argument(
        "--yes",
        action="store_true",
        required=True,
        help="Confirm deletion of the collection",
    )

    args = parser.parse_args()

    if args.command == "csv":
        ingest_csv(csv_path=args.path, force=args.force)

    elif args.command == "files":
        ingest_files(folder=args.folder, module=args.module, issue_type=args.issue_type)

    elif args.command == "stats":
        stats = get_collection_stats(detailed=True)
        print(json.dumps(stats, indent=2))

    elif args.command == "verify":
        run_verify(sample_size=args.sample_size, k=args.k)

    elif args.command == "reset":
        if not args.yes:
            print("Refusing to reset without --yes flag.")
            sys.exit(1)
        reset_collection()
        print("Vector store collection reset successfully.")


if __name__ == "__main__":
    main()
